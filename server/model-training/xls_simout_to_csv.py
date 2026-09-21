"""Convert the simulation output Excel file into the model training CSV.

Usage
-----

    uv run xls_simout_to_csv.py <input.xlsx> <output.csv> [--sheet NAME]

The format of the resulting CSV is documented in ``README.md``.

Sheet
-----

All data is read from the sheet ``Für_Greendelta`` (configurable via
``--sheet``); the name must match exactly, including the ``ü``.  Columns are
addressed by their fixed position (index), never by the header text, and the
header row is always skipped.

Excel input format
------------------

    +------+-----+------------------+---------------------------------------+
    | Col  | Idx | Header           | Meaning                               |
    +------+-----+------------------+---------------------------------------+
    | A    | 0   | GroundSurface    | Ground surface area of the building   |
    |      |     |                  | in m2                                 |
    | B    | 1   | Height           | Building height in m                  |
    | C    | 2   | WeatherStation   | Weather station code (1..15)          |
    | D    | 3   | ConstructionYear | Construction year code (1..6)         |
    | E    | 4   | RoofType         | Roof type code ("1000" or "3100")     |
    | F    | 5   | classification   | Building classification code (1..10)  |
    | G    | 6   | Heat demand      | Annual heat demand in kWh (negative)  |
    | H    | 7   | max. Heat        | Peak load in W (negative)             |
    +------+-----+------------------+---------------------------------------+

Conversion rules
----------------

The meaning of the codes is documented in ``README.md``.

``GroundSurface`` / ``Height`` (columns 0 and 1)
    Copied verbatim; both values must be positive.  ``GroundSurface`` is the
    building footprint (Grundfläche) in m2.

``WeatherStation`` (column 2)
    Copied verbatim; no factor mapping is applied.

``ConstructionYear`` (column 3)
    Copied verbatim, except that an empty cell or ``0`` (year unknown) is
    replaced by the default code ``4`` (1979-1995).  Such substitutions are
    reported as a warning after the conversion.

``RoofType`` (column 4)
    The simulation only knows two roof types, which are mapped to the binary
    flag: ``1000`` (Flachdach) becomes ``1`` and ``3100`` (Satteldach) becomes
    ``0``.  The mapping is deliberately inverted relative to the common
    "1 = pitched" convention, because the previous model gave the flat roof the
    highest roof factor.

``classification`` (column 5)
    Copied verbatim.

``Heat demand`` / ``max. Heat`` (columns 6 and 7)
    The simulation reports both values with a negative sign, so their absolute
    value is used.  The peak load is additionally divided by 1000, because the
    simulation reports it in Watt while the CSV uses kW.

Invalid data
------------

The conversion is atomic: if any data row is invalid, the whole conversion is
aborted and no output file is written.  All errors are collected and reported
together, each with the Excel row number (1-based, including the header) and
the column name.  Fully empty rows (e.g. at the end of a sheet) are skipped.
"""

import argparse
import csv
import math
import sys
from pathlib import Path

from openpyxl import load_workbook

DEFAULT_SHEET_NAME = "Für_Greendelta"

# Column indices.  The Excel sheet and the CSV output use the same column
# order, so a single set of constants addresses both.
COL_GROUND_AREA = 0
COL_HEIGHT = 1
COL_WEATHER_STATION = 2
COL_CONSTRUCTION_YEAR = 3
COL_ROOF_TYPE = 4
COL_BUILDING_TYPE = 5
COL_HEAT_DEMAND = 6
COL_PEAK_LOAD = 7
COL_COUNT = 8

CSV_HEADER = [
    "ground area [m2]",
    "height [m]",
    "weather station [code]",
    "construction year [code]",
    "roof type [1|0]",
    "building type [code]",
    "heat demand [kWh]",
    "peak load [kW]",
]

# Excel roof type code -> (CSV flag, German name)
ROOF_TYPES = {
    "1000": (1, "Flachdach"),
    "3100": (0, "Satteldach"),
}

# The simulation reports the peak load in Watt, the CSV uses kilowatt.
WATT_PER_KILOWATT = 1000

# Valid code ranges (see the module docstring for their meaning).
WEATHER_STATION_MIN, WEATHER_STATION_MAX = 1, 15
CONSTRUCTION_YEAR_MIN, CONSTRUCTION_YEAR_MAX = 1, 6
BUILDING_TYPE_MIN, BUILDING_TYPE_MAX = 1, 10

# Construction year used when the code is unknown (0 or empty).
DEFAULT_CONSTRUCTION_YEAR = 4


class ConversionError(Exception):
    """Raised when the Excel input cannot be converted."""


def _location(row: int, column: str) -> str:
    """Return a human readable location for an error message."""
    return f"Row {row}, column {column!r}"


def _is_empty(value) -> bool:
    """Return whether an Excel cell value counts as empty."""
    return value is None or (isinstance(value, str) and not value.strip())


def _to_float(value, row: int, column: str) -> float:
    """Read a cell as a finite float or raise a conversion error."""
    if _is_empty(value):
        raise ConversionError(f"{_location(row, column)} is empty")
    try:
        number = float(value)
    except (TypeError, ValueError):
        raise ConversionError(
            f"{_location(row, column)} is not a number: {value!r}"
        ) from None
    if not math.isfinite(number):
        raise ConversionError(
            f"{_location(row, column)} is not a finite number: {value!r}"
        )
    return number


def _to_code(
    value, row: int, column: str, low: int, high: int, name: str
) -> int:
    """Read a cell as an integer code within the inclusive range low..high."""
    number = _to_float(value, row, column)
    code = int(number)
    if code != number:
        raise ConversionError(
            f"{_location(row, column)} is not an integer {name} code: {value!r}"
        )
    if not low <= code <= high:
        raise ConversionError(
            f"{_location(row, column)} has an invalid {name} code {code} "
            f"(expected {low}..{high})"
        )
    return code


def _to_positive(value, row: int, column: str) -> float:
    """Read a cell as a strictly positive float."""
    number = _to_float(value, row, column)
    if number <= 0:
        raise ConversionError(
            f"{_location(row, column)} must be positive but is {number}"
        )
    return number


def _to_target(value, row: int, column: str) -> float:
    """Read a heat demand / peak load cell as a positive magnitude.

    The simulation reports both targets with a negative sign, so we take the
    absolute value.  A value of zero indicates missing simulation output and is
    rejected.
    """
    number = abs(_to_float(value, row, column))
    if number == 0:
        raise ConversionError(
            f"{_location(row, column)} is zero (missing simulation output?)"
        )
    return number


def _roof_type_key(value) -> str:
    """Normalize a roof type cell to the code used in ``ROOF_TYPES``.

    Handles numeric cells (1000 / 1000.0) as well as text cells ("1000").
    """
    if _is_empty(value):
        return ""
    if isinstance(value, bool):
        return str(value)
    text = str(value).strip()
    try:
        number = float(text)
    except ValueError:
        return text
    return str(int(number)) if number.is_integer() else text


def _to_roof_type(value, row: int) -> int:
    """Map an Excel roof type to the binary CSV flag."""
    key = _roof_type_key(value)
    if key not in ROOF_TYPES:
        expected = " or ".join(sorted(ROOF_TYPES))
        raise ConversionError(
            f"{_location(row, 'RoofType')} has unknown roof type {value!r} "
            f"(expected {expected})"
        )
    return ROOF_TYPES[key][0]


def _convert_row(values: list, row: int) -> tuple[list, bool]:
    """Convert one Excel row into a CSV row.

    Returns the CSV row and whether the default construction year was used.
    """
    ground_area = _to_positive(values[COL_GROUND_AREA], row, "GroundSurface")
    height = _to_positive(values[COL_HEIGHT], row, "Height")
    weather_station = _to_code(
        values[COL_WEATHER_STATION],
        row,
        "WeatherStation",
        WEATHER_STATION_MIN,
        WEATHER_STATION_MAX,
        "weather station",
    )

    year_value = values[COL_CONSTRUCTION_YEAR]
    used_default = (
        _is_empty(year_value)
        or _to_float(year_value, row, "ConstructionYear") == 0
    )
    construction_year = (
        DEFAULT_CONSTRUCTION_YEAR
        if used_default
        else _to_code(
            year_value,
            row,
            "ConstructionYear",
            CONSTRUCTION_YEAR_MIN,
            CONSTRUCTION_YEAR_MAX,
            "construction year",
        )
    )

    roof_type = _to_roof_type(values[COL_ROOF_TYPE], row)
    building_type = _to_code(
        values[COL_BUILDING_TYPE],
        row,
        "classification",
        BUILDING_TYPE_MIN,
        BUILDING_TYPE_MAX,
        "building type",
    )
    heat_demand = _to_target(values[COL_HEAT_DEMAND], row, "Heat demand")
    peak_load = (
        _to_target(values[COL_PEAK_LOAD], row, "max. Heat") / WATT_PER_KILOWATT
    )

    csv_row = [
        ground_area,
        height,
        weather_station,
        construction_year,
        roof_type,
        building_type,
        heat_demand,
        peak_load,
    ]
    return csv_row, used_default


def _read_rows(path: Path, sheet_name: str) -> list[tuple[int, list]]:
    """Read the data rows (without the header) from the Excel sheet."""
    if not path.is_file():
        raise ConversionError(f"Input file does not exist: {path}")

    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as e:
        raise ConversionError(f"Failed to open Excel file {path}: {e}") from e

    try:
        if sheet_name not in workbook.sheetnames:
            available = ", ".join(workbook.sheetnames)
            raise ConversionError(
                f"Sheet {sheet_name!r} not found in {path.name} "
                f"(available sheets: {available})"
            )

        rows: list[tuple[int, list]] = []
        for index, raw in enumerate(
            workbook[sheet_name].iter_rows(values_only=True)
        ):
            if index == 0:
                continue  # skip the header row
            values = list(raw[:COL_COUNT])
            if len(values) < COL_COUNT:
                values.extend([None] * (COL_COUNT - len(values)))
            if all(_is_empty(v) for v in values):
                continue  # skip fully empty rows
            rows.append((index + 1, values))  # 1-based, matching Excel
        return rows
    finally:
        workbook.close()


def _convert_rows(rows: list[tuple[int, list]]) -> tuple[list[list], int]:
    """Convert all rows, collecting errors instead of failing on the first."""
    data: list[list] = []
    errors: list[str] = []
    defaulted = 0

    for row, values in rows:
        try:
            csv_row, used_default = _convert_row(values, row)
        except ConversionError as e:
            errors.append(str(e))
            continue
        data.append(csv_row)
        if used_default:
            defaulted += 1

    if errors:
        raise ConversionError("\n".join(errors))
    return data, defaulted


def _write_csv(path: Path, data: list[list]) -> None:
    """Write the converted rows to the CSV output file."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(CSV_HEADER)
        writer.writerows(data)


def convert(
    input_file: Path, output_file: Path, sheet_name: str = DEFAULT_SHEET_NAME
) -> tuple[int, int]:
    """Convert the Excel input into the training CSV.

    Returns the number of written rows and the number of rows where the default
    construction year was used.
    """
    rows = _read_rows(input_file, sheet_name)
    if not rows:
        raise ConversionError(
            f"Sheet {sheet_name!r} in {input_file.name} contains no data rows"
        )
    data, defaulted = _convert_rows(rows)
    _write_csv(output_file, data)
    return len(data), defaulted


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Convert the simulation output Excel file into the model "
        "training CSV.",
    )
    parser.add_argument(
        "input", type=Path, help="the simulation output .xlsx file"
    )
    parser.add_argument("output", type=Path, help="the CSV file to write")
    parser.add_argument(
        "--sheet",
        default=DEFAULT_SHEET_NAME,
        help=f"the sheet to read (default: {DEFAULT_SHEET_NAME})",
    )
    args = parser.parse_args()

    try:
        count, defaulted = convert(args.input, args.output, args.sheet)
    except ConversionError as e:
        print(f"Conversion failed:\n{e}", file=sys.stderr)
        return 1

    print(f"Wrote {count} rows to {args.output}")
    if defaulted:
        print(
            f"Warning: construction year was unknown in {defaulted} row(s); "
            f"used default code {DEFAULT_CONSTRUCTION_YEAR} (1979-1995)."
        )
    return 0


if __name__ == "__main__":
    sys.exit(main())
